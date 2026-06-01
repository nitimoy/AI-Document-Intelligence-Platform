"""Document parsing service with factory-based parser selection."""

from __future__ import annotations

import shutil
import subprocess
from abc import ABC, abstractmethod
from pathlib import Path

import fitz
from docx import Document

from app.models.parsed_document import ParsedDocument
from app.utils.exceptions import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from app.utils.logger import get_logger

MIN_CONTENT_LENGTH = 10
MIN_WORD_COUNT = 2
LEGACY_DOC_SUPPORTED = shutil.which("textutil") is not None


class BaseParser(ABC):
    """Abstract parser contract for document format-specific extraction."""

    @abstractmethod
    def parse(self, file_path: Path) -> str:
        """Extract text content from a file path."""


class PDFParser(BaseParser):
    """Extract text content from PDF files using PyMuPDF."""

    def parse(self, file_path: Path) -> str:
        """Extract text from all pages while preserving page order."""

        try:
            with fitz.open(file_path) as pdf:
                if pdf.is_encrypted:
                    # Attempt empty-password auth first for lightly protected files.
                    if not pdf.authenticate(""):
                        raise DocumentParsingError(
                            "PDF is encrypted and cannot be read"
                        )

                page_text = [page.get_text("text") for page in pdf]
                return "\n".join(page_text).strip()
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(f"Failed to parse PDF: {exc}") from exc


class DOCXParser(BaseParser):
    """Extract text content from DOCX files using python-docx."""

    def parse(self, file_path: Path) -> str:
        """Extract paragraph text in source order while ignoring formatting."""

        try:
            doc = Document(str(file_path))
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            return "\n".join(paragraphs).strip()
        except Exception as exc:
            raise DocumentParsingError(f"Failed to parse DOCX: {exc}") from exc


class DOCParser(BaseParser):
    """Extract text content from legacy .doc files via macOS textutil."""

    def parse(self, file_path: Path) -> str:
        """Convert DOC to text via textutil stdout and return extracted text."""

        try:
            process = subprocess.run(
                [
                    "textutil",
                    "-convert",
                    "txt",
                    "-stdout",
                    str(file_path),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            return process.stdout.strip()
        except FileNotFoundError as exc:
            raise DocumentParsingError(
                "textutil is required to parse .doc files on this system"
            ) from exc
        except subprocess.CalledProcessError as exc:
            stderr = exc.stderr.strip() if exc.stderr else "unknown error"
            raise DocumentParsingError(f"Failed to parse DOC: {stderr}") from exc
        except Exception as exc:
            raise DocumentParsingError(f"Failed to parse DOC: {exc}") from exc


class TXTParser(BaseParser):
    """Extract text content from plain-text files with encoding fallback."""

    def parse(self, file_path: Path) -> str:
        """Try UTF-8 first, then latin-1 for legacy text documents."""

        try:
            return file_path.read_text(encoding="utf-8").strip()
        except UnicodeDecodeError:
            try:
                return file_path.read_text(encoding="latin-1").strip()
            except Exception as exc:
                raise DocumentParsingError(
                    f"Failed to parse TXT with latin-1 fallback: {exc}"
                ) from exc
        except Exception as exc:
            raise DocumentParsingError(f"Failed to parse TXT: {exc}") from exc


class DocumentParserFactory:
    """Factory for resolving parser implementations by file extension."""

    _parsers: dict[str, BaseParser] = {
        ".pdf": PDFParser(),
        ".docx": DOCXParser(),
        ".txt": TXTParser(),
    }

    if LEGACY_DOC_SUPPORTED:
        _parsers[".doc"] = DOCParser()

    @classmethod
    def get_parser(cls, file_path: Path) -> BaseParser:
        """Return a parser instance for the file extension."""

        parser = cls._parsers.get(file_path.suffix.lower())
        if parser is None:
            raise UnsupportedFileTypeError(f"Unsupported file type: {file_path.suffix}")
        return parser


class ParserService:
    """Coordinates parser selection, extraction, and content validation."""

    def __init__(self) -> None:
        self._logger = get_logger(__name__)

    def validate_content(self, content: str) -> None:
        """Validate extracted content quality and raise meaningful errors."""

        stripped = content.strip()
        if not stripped:
            raise EmptyDocumentError("Extracted content is empty")
        if len(stripped) < MIN_CONTENT_LENGTH:
            raise EmptyDocumentError(
                "Extracted content is too short "
                f"(minimum {MIN_CONTENT_LENGTH} characters)"
            )

        word_count = len(stripped.split())
        if word_count < MIN_WORD_COUNT:
            raise EmptyDocumentError(
                f"Extracted content has too few words (minimum {MIN_WORD_COUNT})"
            )

    def parse_file(self, file_path: str) -> ParsedDocument:
        """Parse one file and return extraction metadata and content."""

        path = Path(file_path)
        file_name = path.name
        file_type = path.suffix.lower().lstrip(".")

        self._logger.info("File received for parsing: %s", path)

        try:
            parser = DocumentParserFactory.get_parser(path)
            self._logger.info(
                "Parser selected: %s for %s", parser.__class__.__name__, file_name
            )
            self._logger.info("Extraction started for %s", file_name)

            content = parser.parse(path)
            self.validate_content(content)

            word_count = len(content.split())
            character_count = len(content)

            self._logger.info("Extraction completed for %s", file_name)
            self._logger.info("Word count for %s: %s", file_name, word_count)
            self._logger.info("Character count for %s: %s", file_name, character_count)

            return ParsedDocument(
                file_name=file_name,
                file_type=file_type,
                file_path=str(path),
                content=content,
                word_count=word_count,
                character_count=character_count,
                extraction_success=True,
                error_message=None,
            )
        except (
            UnsupportedFileTypeError,
            DocumentParsingError,
            EmptyDocumentError,
        ) as exc:
            self._logger.error("Extraction failed for %s: %s", file_name, exc)
            if isinstance(exc, EmptyDocumentError):
                self._logger.error("Validation failed for %s: %s", file_name, exc)
            return ParsedDocument(
                file_name=file_name,
                file_type=file_type,
                file_path=str(path),
                content="",
                word_count=0,
                character_count=0,
                extraction_success=False,
                error_message=str(exc),
            )

    def parse_multiple_files(self, file_paths: list[str]) -> list[ParsedDocument]:
        """Parse multiple files and return individual parsed document results."""

        return [self.parse_file(file_path) for file_path in file_paths]
